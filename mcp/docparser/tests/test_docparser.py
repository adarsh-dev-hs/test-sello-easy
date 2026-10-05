import functools
import http.server
import io
import os
import threading

import pytest
from fastmcp import Client

import server
from server import mcp

FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"


def _make_files(root):
    # PDF with text (page 1) and an image-only page (page 2) to exercise OCR fallback
    import fitz
    from PIL import Image, ImageDraw, ImageFont

    doc = fitz.open()
    p = doc.new_page()
    p.insert_text((72, 72), "NordWave Networks brochure: private 5G for ports and factories.", fontsize=12)
    img = Image.new("RGB", (1200, 400), "white")
    d = ImageDraw.Draw(img)
    font = ImageFont.truetype(FONT, 64) if os.path.exists(FONT) else ImageFont.load_default()
    d.text((40, 150), "Scanned Page Latency", fill="black", font=font)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    p2 = doc.new_page()
    p2.insert_image(p2.rect, stream=buf.getvalue())
    (root / "brochure.pdf").write_bytes(doc.tobytes())

    # image for OCR
    img2 = Image.new("RGB", (1200, 300), "white")
    ImageDraw.Draw(img2).text((40, 100), "NW-Edge Core Latency 10 ms", fill="black", font=font)
    img2.save(root / "arch.png")

    # docx
    import docx

    dd = docx.Document()
    dd.add_heading("Security Overview", 1)
    dd.add_paragraph("LedgerLeaf is SOC 2 Type II certified.")
    t = dd.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Control", "Status"
    t.cell(1, 0).text, t.cell(1, 1).text = "Encryption at rest", "AES-256"
    dd.save(root / "sec.docx")

    # xlsx with two sheets
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pricing tiers"
    ws.append(["Tier", "Price"])
    ws.append(["Starter", 1000])
    ws2 = wb.create_sheet("Reference customers")
    ws2.append(["Customer", "Industry"])
    ws2.append(["Fjordhavn Terminal", "Ports"])
    wb.save(root / "pricing.xlsx")

    (root / "data.csv").write_text("a,b\n1,2\n3,4\n")
    (root / "notes.md").write_text("# Notes\nHello world")
    (root / "noext").write_bytes((root / "sec.docx").read_bytes())
    (root / "clip.mp3").write_bytes(b"ID3" + b"\x00" * 100)


@pytest.fixture(scope="module")
def base_url(tmp_path_factory):
    root = tmp_path_factory.mktemp("files")
    _make_files(root)
    handler = functools.partial(http.server.SimpleHTTPRequestHandler, directory=str(root))
    handler.log_message = lambda *a, **k: None
    httpd = http.server.ThreadingHTTPServer(("127.0.0.1", 0), handler)
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{httpd.server_address[1]}"
    httpd.shutdown()


async def _parse(base_url, name, **kw):
    async with Client(mcp) as c:
        res = await c.call_tool("parse_document", {"file_url": f"{base_url}/{name}", **kw})
    d = res.data
    assert set(d) >= {"text", "tables", "pages", "meta"}
    return d


async def test_pdf_with_ocr_page(base_url):
    d = await _parse(base_url, "brochure.pdf")
    assert d["pages"] == 2
    assert "private 5G" in d["text"]
    assert d["meta"]["ocr_pages"] == [2]
    assert "Latency" in d["text"]


async def test_image_ocr(base_url):
    d = await _parse(base_url, "arch.png")
    assert "Latency" in d["text"] and "Edge" in d["text"]
    assert d["meta"]["kind"] == "image"


async def test_docx(base_url):
    d = await _parse(base_url, "sec.docx")
    assert "SOC 2" in d["text"]
    assert d["tables"] and "AES-256" in d["tables"][0]
    assert "AES-256" in d["text"]


async def test_docx_detected_by_magic(base_url):
    d = await _parse(base_url, "noext", mime_type="application/octet-stream")
    assert d["meta"]["kind"] == "docx"


async def test_xlsx(base_url):
    d = await _parse(base_url, "pricing.xlsx")
    assert len(d["tables"]) == 2
    assert "Fjordhavn Terminal" in d["text"]
    assert d["meta"]["sheets"] == ["Pricing tiers", "Reference customers"]


async def test_csv_and_text(base_url):
    d = await _parse(base_url, "data.csv")
    assert d["tables"] and "| a | b |" in d["tables"][0]
    d = await _parse(base_url, "notes.md")
    assert "Hello world" in d["text"]


async def test_audio_without_key(base_url, monkeypatch):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    d = await _parse(base_url, "clip.mp3")
    assert d["text"] == "" and "warning" in d["meta"]


async def test_download_error():
    async with Client(mcp) as c:
        res = await c.call_tool("parse_document", {"file_url": "http://127.0.0.1:1/missing.pdf"})
    assert res.data["text"] == "" and "error" in res.data["meta"]


def test_detect_kind():
    assert server.detect_kind(b"%PDF-1.7", "", "", "") == "pdf"
    assert server.detect_kind(b"hello", "text/csv", "", "") == "csv"
    assert server.detect_kind(b"hello", "", "x.md", "") == "text"
