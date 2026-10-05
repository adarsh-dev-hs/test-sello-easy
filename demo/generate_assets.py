#!/usr/bin/env python3
"""Generate the SelloQ demo documents for the two fictional sellers.

Usage:  python demo/generate_assets.py [OUTPUT_DIR]        (default: demo/docs next to this script)

Writes:
  nordwave/nordwave_brochure.pdf, nordwave/nordwave_pricing.xlsx, nordwave/nordwave_architecture.png
  ledgerleaf/ledgerleaf_case_study.pdf, ledgerleaf/ledgerleaf_roi_calculator.xlsx,
  ledgerleaf/ledgerleaf_dashboard.png, ledgerleaf/ledgerleaf_security_overview.docx

Deps: reportlab, openpyxl, Pillow, python-docx (see demo/requirements.txt). All content is fictional.
"""
from __future__ import annotations

import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/Library/Fonts/Arial.ttf",
    "C:/Windows/Fonts/arialbd.ttf",
    "DejaVuSans-Bold.ttf",
]


def font(size: int):
    for path in FONT_CANDIDATES:
        try:
            return ImageFont.truetype(path, size)
        except Exception:
            continue
    try:
        return ImageFont.load_default(size=size)  # Pillow >= 10.1
    except TypeError:
        return ImageFont.load_default()


# ----------------------------------------------------------------------------- PDF helpers
def _styles():
    ss = getSampleStyleSheet()
    ss.add(ParagraphStyle("Body", parent=ss["BodyText"], fontSize=10.5, leading=15, spaceAfter=6))
    ss.add(ParagraphStyle("H1x", parent=ss["Heading1"], fontSize=20, spaceAfter=10))
    ss.add(ParagraphStyle("H2x", parent=ss["Heading2"], fontSize=14, spaceBefore=10, spaceAfter=6))
    return ss


def _table(data, col_widths=None, header_color="#0b2a4a"):
    t = Table(data, colWidths=col_widths, repeatRows=1)
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(header_color)),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9.5),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#c9d3df")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f6fa")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    return t


# ----------------------------------------------------------------------------- NordWave
def nordwave_brochure(path: str):
    ss = _styles()
    P = lambda t, s="Body": Paragraph(t, ss[s])  # noqa: E731
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm,
                            bottomMargin=2 * cm, title="NordWave Networks Brochure", author="NordWave Networks")
    story = [
        P("NordWave Networks — Private 5G for Industry", "H1x"),
        P("<b>Keep cranes, robots, AGVs and sensors connected — with a private network you own and control.</b>"),
        P("Industrial operations are automating fast, but most plants, ports and mines still depend on enterprise "
          "Wi-Fi that was designed for offices. AGVs stop during roaming, crane telemetry drops out behind steel "
          "structures, and Industry 4.0 programmes stall at the pilot stage. NordWave Networks delivers a complete, "
          "managed private 5G platform for operational technology (OT) that removes the network as the bottleneck."),
        P("The platform", "H2x"),
        P("<b>NW-Edge private 5G core</b> — an on-premise 5G standalone core with network slicing, SIM/eSIM identity "
          "and local breakout for user-plane latency below 10 ms. Redundant deployment provides 99.999% availability, "
          "and an optional edge compute module hosts vision, safety and digital-twin workloads."),
        P("<b>NW-Radio industrial small cells</b> — IP67, -40°C to +60°C, ATEX Zone 2 variants. One cell covers up to "
          "25,000 m² indoors or 1 km outdoors and replaces 6–10 Wi-Fi access points. Models: NW-Radio 400 (indoor), "
          "NW-Radio 800 (outdoor yards, quays, open pits) and NW-Radio X (hazardous areas)."),
        P("<b>NW-Ops analytics</b> — live dashboards for latency, packet loss and coverage per zone, alerts before an "
          "AGV link degrades, and correlation of network events with production stops. Backed by NordWave's 24/7 "
          "network operations centre."),
        P("Who we serve", "H2x"),
        P("Ports and terminals, manufacturing (automotive, steel, chemicals, paper, food), mining and utilities with "
          "1,000+ employees in Europe and North America. Typical buyers are VP Operations, Plant Directors, Heads of "
          "OT, CIOs and Chief Digital Officers."),
        P("Key benefits", "H2x"),
        P("• Up to 98% fewer connectivity-related AGV stops<br/>• Zero-loss handover at up to 60 km/h<br/>"
          "• 6–10x fewer radios than Wi-Fi<br/>• Works with local spectrum licences (3.7–3.8 GHz EU) and CBRS (US)<br/>"
          "• Per-site subscription — no large upfront capex"),
        PageBreak(),
        P("Technical specifications", "H1x"),
        _table(
            [
                ["Parameter", "NW-Edge / NW-Radio specification"],
                ["End-to-end latency (user plane)", "< 10 ms (typical 6 ms)"],
                ["Availability", "99.999% with redundant NW-Edge pair"],
                ["Devices per site", "up to 20,000"],
                ["Handover interruption", "0 ms (make-before-break)"],
                ["Spectrum bands", "n77 / n78 (3.3–4.2 GHz), n48 CBRS, LTE B42/B43"],
                ["Radio coverage", "25,000 m² indoor / 1 km outdoor per cell"],
                ["Environmental", "IP67, -40°C to +60°C, ATEX Zone 2 (NW-Radio X)"],
                ["Security", "SIM/eSIM auth, IEC 62443 aligned, on-premise data"],
                ["Integrations", "OPC UA, Modbus TCP gateways, SCADA, MES, Active Directory"],
            ],
            col_widths=[6.5 * cm, 10.5 * cm],
        ),
        Spacer(1, 14),
        P("Deployment model", "H2x"),
        _table(
            [
                ["Phase", "Duration", "What happens"],
                ["Site survey & radio plan", "2 weeks", "RF survey, spectrum check, success metrics agreed"],
                ["Pilot", "8–12 weeks", "One hall, yard or mine level; AGV and telemetry KPIs measured"],
                ["Scale-out", "3–9 months", "Full-site coverage, OT integration, NW-Ops rollout"],
                ["Managed operations", "Ongoing", "24/7 NOC, SLA-backed, quarterly optimisation reviews"],
            ],
            col_widths=[4.5 * cm, 3 * cm, 9.5 * cm],
        ),
        Spacer(1, 14),
        P("Customer results", "H2x"),
        P("Fjordhavn Container Terminal (Norway): +40% crane-move productivity. Elbtal Automotive (Germany): 98% fewer "
          "AGV connectivity stops. Baltic Forge (Lithuania): EUR 2.3m per year in avoided downtime. Copper Ridge Mine "
          "(USA): 99.95% network availability for 40 autonomous haul trucks."),
        P("Contact: sales@nordwave.example · www.nordwave.example (fictional demo company)"),
    ]
    doc.build(story)


def nordwave_pricing(path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Pricing tiers"
    rows = [
        ["Tier", "Monthly price (EUR)", "Small cells included", "Coverage", "Devices", "Support", "Typical use"],
        ["Site Starter", 4900, 4, "up to 100,000 m²", 500, "Business hours + NOC monitoring", "Pilot hall, yard or mine level"],
        ["Industrial", 12500, 12, "up to 300,000 m²", 5000, "24/7 NOC, 4h on-site", "Full plant, terminal or mine"],
        ["Enterprise Multi-site", "Custom", "Custom", "Multiple sites", 20000, "24/7 NOC, 2h on-site, TAM", "Groups with 3+ sites"],
        ["Add-on: extra NW-Radio cell", 450, 1, "+25,000 m²", "", "", "Coverage extension"],
        ["Add-on: NW-Edge compute module", 1200, "", "", "", "", "Edge AI / vision workloads"],
        ["One-time: site survey & installation", 18000, "", "", "", "", "Per site, credited on 3-year term"],
    ]
    for r in rows:
        ws.append(r)
    ws.append([])
    ws.append(["Notes: prices exclude VAT; 36-month term; spectrum fees not included. Fictional demo data."])
    rc = wb.create_sheet("Reference customers")
    rc_rows = [
        ["Customer", "Country", "Industry", "Use case", "Result", "Tier"],
        ["Fjordhavn Container Terminal", "Norway", "Ports", "Remote-controlled cranes, straddle carriers", "+40% crane-move productivity", "Industrial"],
        ["Elbtal Automotive", "Germany", "Automotive manufacturing", "120 AGVs in body & assembly shops", "-98% AGV connectivity stops", "Industrial"],
        ["Baltic Forge", "Lithuania", "Steel & forging", "Crane control, thermal cameras", "EUR 2.3m/yr avoided downtime", "Industrial"],
        ["Copper Ridge Mine", "USA", "Mining", "40 autonomous haul trucks", "99.95% network availability", "Enterprise Multi-site"],
        ["Great Lakes Power Cooperative", "USA", "Utilities", "CBRS private LTE for 300 substations", "USD 1.1m/yr saved on leased lines", "Enterprise Multi-site"],
    ]
    for r in rc_rows:
        rc.append(r)
    for sheet in (ws, rc):
        for c in sheet[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="0B2A4A")
            c.alignment = Alignment(wrap_text=True)
        for col in sheet.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            sheet.column_dimensions[col[0].column_letter].width = min(max(12, width + 2), 45)
    wb.save(path)


def _box(d: ImageDraw.ImageDraw, xy, label_lines, fill, outline="#0b2a4a", size=34):
    d.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=4)
    f = font(size)
    x0, y0, x1, y1 = xy
    total_h = sum(d.textbbox((0, 0), ln, font=f)[3] + 10 for ln in label_lines) - 10
    y = y0 + ((y1 - y0) - total_h) / 2
    for ln in label_lines:
        w = d.textbbox((0, 0), ln, font=f)[2]
        d.text((x0 + ((x1 - x0) - w) / 2, y), ln, fill="#111111", font=f)
        y += d.textbbox((0, 0), ln, font=f)[3] + 10


def nordwave_architecture(path: str):
    img = Image.new("RGB", (1200, 800), "white")
    d = ImageDraw.Draw(img)
    d.text((40, 30), "NordWave Private 5G Architecture", fill="#0b2a4a", font=font(48))
    _box(d, (60, 140, 400, 300), ["AGV / Robots", "Sensors"], "#e3f2fd")
    _box(d, (440, 140, 780, 300), ["NW-Radio", "Small Cells"], "#e8f5e9")
    _box(d, (820, 140, 1140, 300), ["NW-Edge Core", "On-Premise"], "#fff3e0")
    _box(d, (440, 400, 780, 560), ["NW-Ops", "Analytics"], "#f3e5f5")
    _box(d, (820, 400, 1140, 560), ["OT Network", "SCADA / MES"], "#fffde7")
    for (a, b) in (((400, 220), (440, 220)), ((780, 220), (820, 220)), ((980, 300), (980, 400)), ((820, 480), (780, 480))):
        d.line([a, b], fill="#0b2a4a", width=6)
    d.text((60, 620), "Latency < 10 ms", fill="#111111", font=font(44))
    d.text((60, 690), "Availability 99.999%", fill="#111111", font=font(44))
    d.text((640, 620), "Network Slicing", fill="#111111", font=font(44))
    d.text((640, 690), "Zero-Loss Handover", fill="#111111", font=font(44))
    img.save(path)


# ----------------------------------------------------------------------------- LedgerLeaf
def ledgerleaf_case_study(path: str):
    ss = _styles()
    P = lambda t, s="Body": Paragraph(t, ss[s])  # noqa: E731
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm,
                            bottomMargin=2 * cm, title="LedgerLeaf Case Study: Northgate Supply Co.", author="LedgerLeaf")
    green = "#1e6b3a"
    story = [
        P("Case Study: Northgate Supply Co. cuts month-end close from 12 to 6 days with LedgerLeaf", "H1x"),
        P("<b>Company:</b> Northgate Supply Co., wholesale distribution, Ohio, USA · 900 employees · 4 warehouses · "
          "3,500 vendor invoices per month · ERP: Oracle NetSuite"),
        P("The challenge", "H2x"),
        P("Northgate's three AP clerks keyed every vendor invoice into NetSuite by hand. Paper invoices arrived at four "
          "warehouses, approvals were collected by email and average approval time was nine days. The company paid "
          "$46,000 in late-payment penalties in one year and missed most early-payment discounts. Month-end close took "
          "12 working days, and with 40% volume growth planned, the Controller was preparing to hire two more AP clerks."),
        P("The solution", "H2x"),
        P("Northgate deployed LedgerLeaf Growth in five weeks: AI invoice capture for email and scanned invoices, "
          "automatic 3-way match against NetSuite purchase orders and item receipts, rules-based approval workflows "
          "for warehouse managers, and real-time two-way NetSuite sync. The supplier portal let vendors check payment "
          "status themselves."),
        P("The results after six months", "H2x"),
        _table(
            [
                ["Metric", "Before LedgerLeaf", "After LedgerLeaf"],
                ["Touchless invoices", "0%", "78%"],
                ["Approval cycle time", "9 days", "1.6 days"],
                ["Month-end close", "12 days", "6 days"],
                ["Cost per invoice", "$12.80", "$2.60"],
                ["Late-payment penalties (annualised)", "$46,000", "$0"],
                ["Early-payment discounts captured", "$9,000", "$71,000"],
                ["AP headcount needed for +40% volume", "5 clerks", "3 clerks (no new hires)"],
            ],
            col_widths=[7 * cm, 5 * cm, 5 * cm],
            header_color=green,
        ),
        Spacer(1, 12),
        P("<i>\"We were about to hire two more AP clerks. Instead we deployed LedgerLeaf and our close went from 12 days "
          "to 6. Our warehouse managers approve invoices from their phones in minutes.\"</i> — Controller, Northgate Supply Co."),
        P("Why LedgerLeaf", "H2x"),
        P("Northgate evaluated a legacy enterprise AP suite and a small-business bill-pay tool. The enterprise suite "
          "required a six-month implementation; the small-business tool lacked 3-way match and multi-warehouse "
          "approvals. LedgerLeaf combined mid-market pricing (no per-user fees) with native NetSuite integration and "
          "went live in weeks. Learn more at www.ledgerleaf.example (fictional demo company)."),
    ]
    doc.build(story)


def ledgerleaf_roi(path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "ROI Calculator"
    ws.append(["LedgerLeaf ROI Calculator (fictional demo)", ""])
    ws.append(["Input", "Value"])
    inputs = [
        ("Invoices per month", 3000),
        ("Manual cost per invoice (USD)", 12.5),
        ("LedgerLeaf cost per invoice (USD)", 2.8),
        ("Annual late-payment penalties today (USD)", 40000),
        ("Annual spend eligible for early-pay discount (USD)", 8000000),
        ("Early-pay discount rate", 0.01),
        ("Share of discounts captured with LedgerLeaf", 0.6),
        ("LedgerLeaf annual subscription (USD)", 42000),
        ("Implementation fee (USD)", 12000),
    ]
    for k, v in inputs:
        ws.append([k, v])
    ws.append([])
    # Values are pre-computed (not Excel formulas) so parsers without a calc engine still see numbers.
    v = dict(inputs)
    processing = v["Invoices per month"] * 12 * (v["Manual cost per invoice (USD)"] - v["LedgerLeaf cost per invoice (USD)"])
    penalties = v["Annual late-payment penalties today (USD)"]
    discounts = (v["Annual spend eligible for early-pay discount (USD)"] * v["Early-pay discount rate"]
                 * v["Share of discounts captured with LedgerLeaf"])
    benefit = processing + penalties + discounts
    cost = v["LedgerLeaf annual subscription (USD)"] + v["Implementation fee (USD)"]
    ws.append(["Output", "Value", "How it is calculated"])
    ws.append(["Annual processing savings (USD)", round(processing), "invoices x 12 x (manual cost - LedgerLeaf cost)"])
    ws.append(["Penalties avoided (USD)", penalties, "current annual late-payment penalties"])
    ws.append(["Discounts captured (USD)", round(discounts), "eligible spend x discount rate x share captured"])
    ws.append(["Total annual benefit (USD)", round(benefit), "sum of the three benefits"])
    ws.append(["Year-1 cost (USD)", cost, "subscription + implementation"])
    ws.append(["Year-1 net benefit (USD)", round(benefit - cost), "benefit - cost"])
    ws.append(["Payback (months)", round(cost / (benefit / 12), 1), "cost / monthly benefit"])
    ws.append(["ROI year 1", f"{(benefit - cost) / cost:.0%}", "net benefit / cost"])
    for c in ws[1]:
        c.font = Font(bold=True, size=14, color="1E6B3A")
    for row in (2, 13):
        for c in ws[row]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1E6B3A")
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 50

    pt = wb.create_sheet("Plans")
    for r in [
        ["Plan", "Price per month (USD)", "Invoices per month", "Entities", "Implementation fee (USD)"],
        ["Essentials", 1500, "up to 1,000", 1, 5000],
        ["Growth", 3500, "up to 5,000", "up to 10", 12000],
        ["Enterprise", "Custom", "5,000+", "Unlimited", "Custom"],
    ]:
        pt.append(r)
    bm = wb.create_sheet("Benchmarks")
    for r in [
        ["Metric", "Manual AP (mid-market average)", "With LedgerLeaf"],
        ["Cost per invoice (USD)", 12.4, 2.7],
        ["Approval cycle (days)", 9, 1.8],
        ["Touchless rate", "0%", "75-85%"],
        ["Month-end close (days)", 11, 6],
    ]:
        bm.append(r)
    for sheet in (pt, bm):
        for c in sheet[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1E6B3A")
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 28
    wb.save(path)


def ledgerleaf_dashboard(path: str):
    img = Image.new("RGB", (1200, 800), "white")
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, 1200, 110), fill="#eaf5e4")
    d.text((40, 28), "LedgerLeaf AP Dashboard", fill="#164f2b", font=font(50))
    tiles = [
        ("Invoices", "3,412"),
        ("Touchless Rate", "81%"),
        ("Approval Cycle", "1.7 days"),
        ("Cost per Invoice", "$2.65"),
    ]
    for i, (label, value) in enumerate(tiles):
        x = 40 + i * 285
        d.rounded_rectangle((x, 150, x + 260, 330), radius=16, outline="#1e6b3a", width=4, fill="white")
        d.text((x + 18, 170), label, fill="#222222", font=font(24))
        d.text((x + 18, 235), value, fill="#111111", font=font(44))
    d.text((40, 380), "Exceptions by Type", fill="#164f2b", font=font(36))
    bars = [("Price Mismatch", 46), ("Missing Receipt", 31), ("Duplicate Invoice", 12), ("No PO", 22)]
    for i, (label, n) in enumerate(bars):
        y = 450 + i * 70
        d.text((40, y + 6), label, fill="#111111", font=font(30))
        d.rectangle((420, y + 4, 420 + n * 12, y + 42), fill="#c5e1a5")
        d.text((1060, y + 6), str(n), fill="#111111", font=font(30))
    d.text((40, 740), "Month-End Close: 6 days   Late Fees: $0", fill="#111111", font=font(32))
    img.save(path)


def ledgerleaf_security(path: str):
    import docx
    from docx.shared import Pt

    doc = docx.Document()
    doc.core_properties.title = "LedgerLeaf Security Overview"
    doc.add_heading("LedgerLeaf Security Overview", 0)
    doc.add_paragraph(
        "This document summarises how LedgerLeaf protects customer financial data. LedgerLeaf is an accounts payable "
        "automation platform for mid-market finance teams in the United States, Canada and the United Kingdom. "
        "(Fictional demo company.)"
    )
    doc.add_heading("Certifications and compliance", 1)
    for t in [
        "SOC 2 Type II report issued annually by an independent auditor",
        "ISO/IEC 27001 certified information security management system",
        "GDPR and UK GDPR compliant; Data Processing Agreement available",
        "Supports customer SOX controls with full approval audit trails",
    ]:
        doc.add_paragraph(t, style="List Bullet")
    doc.add_heading("Data protection", 1)
    doc.add_paragraph(
        "All customer data is encrypted at rest with AES-256 and in transit with TLS 1.3. Encryption keys are managed in "
        "a cloud key management service with automatic rotation. Customers can choose data residency in the US, Canada "
        "or the UK. Invoice images and extracted data are retained according to customer-configured policies."
    )
    doc.add_heading("Access control", 1)
    doc.add_paragraph(
        "LedgerLeaf supports SAML single sign-on (Okta, Entra ID, Google Workspace), enforced MFA, role-based access "
        "control and segregation of duties between invoice entry, approval and payment release. Every action is logged "
        "with user, timestamp and IP address."
    )
    doc.add_heading("Controls summary", 1)
    rows = [
        ("Control", "Implementation"),
        ("Encryption at rest", "AES-256"),
        ("Encryption in transit", "TLS 1.3"),
        ("Authentication", "SAML SSO + MFA"),
        ("Penetration testing", "Annual third-party test + continuous scanning"),
        ("Backups", "Daily, encrypted, 35-day retention, cross-region"),
        ("Availability SLA", "99.9% (Enterprise: 99.95%)"),
        ("Incident response", "24/7 on-call, customer notification within 72 hours"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    for i, (a, b) in enumerate(rows):
        table.cell(i, 0).text = a
        table.cell(i, 1).text = b
        if i == 0:
            for c in table.rows[0].cells:
                for r in c.paragraphs[0].runs:
                    r.font.bold = True
                    r.font.size = Pt(11)
    doc.add_heading("Payments security", 1)
    doc.add_paragraph(
        "Vendor bank-detail changes require dual approval and out-of-band verification, protecting against invoice "
        "fraud and business email compromise. Payment files are signed and transmitted to banks over SFTP."
    )
    doc.add_paragraph("Contact: security@ledgerleaf.example")
    doc.save(path)


# ----------------------------------------------------------------------------- main
def main(out_dir: str) -> list[str]:
    nw = os.path.join(out_dir, "nordwave")
    ll = os.path.join(out_dir, "ledgerleaf")
    os.makedirs(nw, exist_ok=True)
    os.makedirs(ll, exist_ok=True)
    jobs = [
        (nordwave_brochure, os.path.join(nw, "nordwave_brochure.pdf")),
        (nordwave_pricing, os.path.join(nw, "nordwave_pricing.xlsx")),
        (nordwave_architecture, os.path.join(nw, "nordwave_architecture.png")),
        (ledgerleaf_case_study, os.path.join(ll, "ledgerleaf_case_study.pdf")),
        (ledgerleaf_roi, os.path.join(ll, "ledgerleaf_roi_calculator.xlsx")),
        (ledgerleaf_dashboard, os.path.join(ll, "ledgerleaf_dashboard.png")),
        (ledgerleaf_security, os.path.join(ll, "ledgerleaf_security_overview.docx")),
    ]
    written = []
    for fn, path in jobs:
        fn(path)
        written.append(path)
        print("wrote", path)
    return written


if __name__ == "__main__":
    default = os.path.join(os.path.dirname(os.path.abspath(__file__)), "docs")
    main(sys.argv[1] if len(sys.argv) > 1 else default)
